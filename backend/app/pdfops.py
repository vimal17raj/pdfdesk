import base64
import difflib
import io
import json
import os
import re
import secrets
import shutil
import subprocess
import time
import zipfile
from collections import Counter
from contextvars import ContextVar
from datetime import date
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import pymupdf as fitz
from fastapi import HTTPException

from . import config


# ---------------------------------------------------------------- basics
# When set, edit functions modify the document they were given in place and return None; the caller
# (store.edit) then saves only the changes. Otherwise they work on bytes and return the new file's bytes.
INPLACE: ContextVar[bool] = ContextVar("pdfdesk_inplace", default=False)


def _open(data) -> fitz.Document:
    """Accepts an already-open Document (used as is), a file path (opened lazily from disk) or PDF bytes."""
    if isinstance(data, fitz.Document):
        return data
    try:
        doc = fitz.open(str(data), filetype="pdf") if isinstance(data, (str, Path)) else fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise HTTPException(400, "Not a valid PDF file")
    if doc.needs_pass:
        raise HTTPException(400, "PASSWORD_REQUIRED")
    return doc


def _release(doc: fitz.Document):
    """Close a document we opened ourselves, but never the shared, cached one."""
    if not getattr(doc, "_pd_shared", False):
        doc.close()


def _out(doc: fitz.Document, **kw):
    if INPLACE.get():
        return None
    # garbage=1 drops unused objects but keeps object numbers stable, so a selected annotation keeps its id across edits
    return doc.tobytes(garbage=1, deflate=True, **kw)


def _page(doc: fitz.Document, n: int) -> fitz.Page:
    if not 0 <= n < doc.page_count:
        raise HTTPException(404, "No such page")
    return doc[n]


# Pages can carry a /Rotate flag. The browser works in the rotated ("visual") space,
# PDF content lives in the unrotated space; these two helpers convert between them.
def _to_page(pg, r) -> fitz.Rect:
    return (fitz.Rect(r) * pg.derotation_matrix).normalize()


def _pt_to_page(pg, p) -> fitz.Point:
    return fitz.Point(p) * pg.derotation_matrix


def _to_view(pg, r) -> list:
    return list((fitz.Rect(r) * pg.rotation_matrix).normalize())


def unlock(data: bytes, password: str | None) -> bytes:
    """Return the PDF unchanged, or decrypted when it needs (and is given) a password."""
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise HTTPException(400, "Not a valid PDF file")
    if not doc.needs_pass:
        doc.close()
        return data
    if not password:
        raise HTTPException(400, "PASSWORD_REQUIRED")
    if not doc.authenticate(password):
        raise HTTPException(400, "Wrong password")
    out = doc.tobytes(encryption=fitz.PDF_ENCRYPT_NONE, garbage=3, deflate=True)
    doc.close()
    return out


def unlock_file(path: Path, password: str | None) -> int:
    """For an uploaded file on disk: check it opens, decrypt it in place if it needs (and is given) a password.
    Returns the page count. Works on any size: nothing is loaded into memory."""
    try:
        doc = fitz.open(str(path), filetype="pdf")  # the upload's temp name has no .pdf ending
    except Exception:
        raise HTTPException(400, "Not a valid PDF file")
    try:
        if doc.needs_pass:
            if not password:
                raise HTTPException(400, "PASSWORD_REQUIRED")
            if not doc.authenticate(password):
                raise HTTPException(400, "Wrong password")
            tmp = path.with_suffix(".dec")
            doc.save(str(tmp), encryption=fitz.PDF_ENCRYPT_NONE, garbage=0, deflate=True)
            n = doc.page_count
            doc.close()
            os.replace(tmp, path)
            return n
        return doc.page_count
    finally:
        if not doc.is_closed:
            doc.close()


def validate(data: bytes) -> int:
    d = _open(data)
    n = d.page_count
    d.close()
    return n


def render(data, page: int, zoom: float) -> bytes:
    """JPEG is several times faster to encode and smaller on the wire than PNG; fine for on-screen pages."""
    doc = _open(data)
    try:
        pix = _page(doc, page).get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        return pix.tobytes("jpeg", jpg_quality=82)
    finally:
        _release(doc)


def page_sizes(data) -> list[dict]:
    doc = _open(data)
    out = [{"w": p.rect.width, "h": p.rect.height, "rotation": p.rotation} for p in (doc[i] for i in range(doc.page_count))]
    _release(doc)
    return out


def _rgb(c) -> tuple:
    if isinstance(c, (list, tuple)) and len(c) == 3:
        return tuple(max(0, min(1, float(x))) for x in c)
    if isinstance(c, int):
        return ((c >> 16 & 255) / 255, (c >> 8 & 255) / 255, (c & 255) / 255)
    return (0, 0, 0)


# ---------------------------------------------------------------- text editing
def text_lines(data: bytes, page: int) -> list[dict]:
    """Editable text lines, bbox in visual PDF points (top-left origin)."""
    doc = _open(data)
    pg = _page(doc, page)
    lines = []
    for b in pg.get_text("dict")["blocks"]:
        if b.get("type") != 0:
            continue
        for ln in b["lines"]:
            spans = [s for s in ln["spans"] if s["text"].strip()]
            if not spans:
                continue
            lines.append(
                {
                    "bbox": _to_view(pg, ln["bbox"]),
                    "text": "".join(s["text"] for s in ln["spans"]),
                    "size": round(spans[0]["size"], 1),
                    "color": spans[0]["color"],
                    "font": spans[0]["font"],
                }
            )
    _release(doc)
    return lines


def edit_text(data: bytes, page: int, bbox, new_text: str, size: float, color) -> bytes:
    """Replace the text inside bbox: blank the old line (images kept), write the new text in its place."""
    doc = _open(data)
    pg = _page(doc, page)
    r = _to_page(pg, bbox)
    pg.add_redact_annot(r, fill=(1, 1, 1))
    pg.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)
    wide = fitz.Rect(r.x0, r.y0 - 1, max(r.x1, pg.rect.x1 - 20), r.y1 + size * 0.4)
    rc = pg.insert_textbox(wide, new_text, fontsize=size, fontname="helv", color=_rgb(color))
    if rc < 0:  # didn't fit: shrink until it does
        for s in (size * 0.9, size * 0.8, size * 0.7, size * 0.6):
            if pg.insert_textbox(wide, new_text, fontsize=s, fontname="helv", color=_rgb(color)) >= 0:
                break
    return _out(doc)


# ---------------------------------------------------------------- annotations & comments
_HIDDEN_ANNOTS = {"Link", "Widget", "Popup"}


def annotate(data: bytes, page: int, a: dict, author: str = "") -> bytes:
    doc = _open(data)
    pg = _page(doc, page)
    kind = a["type"]
    color = _rgb(a.get("color", [1, 0.9, 0]))
    r = _to_page(pg, a["rect"]) if a.get("rect") else None
    ann = None
    if kind in ("text", "highlight", "underline", "strike", "rect", "ellipse", "redact", "note") and r is None:
        raise HTTPException(400, "rect required")
    if kind == "text":
        ann = pg.add_freetext_annot(
            r, a.get("text", ""), fontsize=a.get("size", 12), text_color=_rgb(a.get("color", [0, 0, 0])), fontname="helv"
        )
        ann.update()
    elif kind == "highlight":
        ann = pg.add_highlight_annot(r)
        ann.set_colors(stroke=color)
    elif kind == "underline":
        ann = pg.add_underline_annot(r)
    elif kind == "strike":
        ann = pg.add_strikeout_annot(r)
    elif kind in ("rect", "ellipse"):
        ann = pg.add_rect_annot(r) if kind == "rect" else pg.add_circle_annot(r)
        ann.set_colors(stroke=color)
        ann.set_border(width=a.get("width", 2))
        ann.update()
    elif kind == "ink":
        strokes = [[tuple(_pt_to_page(pg, (x, y))) for x, y in s] for s in a["strokes"]]
        ann = pg.add_ink_annot(strokes)
        ann.set_colors(stroke=color)
        ann.set_border(width=a.get("width", 2))
        ann.update()
    elif kind == "note":
        ann = pg.add_text_annot(r.tl, a.get("text", ""))
    elif kind == "redact":  # permanent: removes the text/pixels underneath
        pg.add_redact_annot(r, fill=(0, 0, 0))
        pg.apply_redactions()
    else:
        raise HTTPException(400, f"Unknown annotation type {kind}")
    if ann is not None and kind != "text":
        ann.set_info(title=author[:60], content=a.get("text", ""))
        ann.update()
    elif ann is not None:
        ann.set_info(title=author[:60])
    return _out(doc)


def _hex(c) -> str:
    if not c or len(c) != 3:
        return "#000000"
    return "#%02x%02x%02x" % tuple(int(round(v * 255)) for v in c)


def list_annots(data: bytes, page: int) -> list[dict]:
    doc = _open(data)
    pg = _page(doc, page)
    out = [
        {"xref": an.xref, "type": an.type[1], "rect": _to_view(pg, an.rect), "text": an.info.get("content", ""),
         "author": an.info.get("title", ""), "color": _hex(an.colors.get("stroke"))}
        for an in pg.annots() or []
        if an.type[1] not in _HIDDEN_ANNOTS
    ]
    _release(doc)
    return out


def all_comments(data: bytes) -> list[dict]:
    doc = _open(data)
    out = []
    for i, pg in enumerate(doc):
        for an in pg.annots() or []:
            t = an.type[1]
            if t in _HIDDEN_ANNOTS:
                continue
            out.append({"page": i, "xref": an.xref, "type": t, "text": an.info.get("content", ""),
                        "author": an.info.get("title", ""), "rect": _to_view(pg, an.rect)})
    _release(doc)
    return out


def delete_annot(data: bytes, page: int, xref: int) -> bytes:
    doc = _open(data)
    pg = _page(doc, page)
    for an in pg.annots() or []:
        if an.xref == xref:
            pg.delete_annot(an)
            return _out(doc)
    raise HTTPException(404, "Annotation not found")


# ---------------------------------------------------------------- page organisation
def pages_op(data: bytes, op: str, pages: list[int], angle: int = 90, order: list[int] | None = None) -> bytes:
    doc = _open(data)
    n = doc.page_count
    if any(not 0 <= p < n for p in pages) or (order and sorted(order) != list(range(n))):
        raise HTTPException(400, "Invalid page numbers")
    if op == "rotate":
        for p in pages:
            doc[p].set_rotation((doc[p].rotation + angle) % 360)
    elif op == "delete":
        if len(set(pages)) >= n:
            raise HTTPException(400, "Cannot delete all pages")
        doc.delete_pages(sorted(set(pages)))
    elif op == "reorder":
        doc.select(order)
    elif op == "reverse":
        doc.select(list(range(n))[::-1])
    elif op == "duplicate":
        for p in sorted(set(pages), reverse=True):
            doc.fullcopy_page(p, p + 1)
    elif op == "blank":
        for p in sorted(set(pages), reverse=True):
            ref = doc[p].rect
            doc.new_page(p + 1, width=ref.width, height=ref.height)
    else:
        raise HTTPException(400, "Unknown page operation")
    return _out(doc)


def extract(data: bytes, pages: list[int]) -> bytes:
    doc = _open(data)
    if not pages or any(not 0 <= p < doc.page_count for p in pages):
        raise HTTPException(400, "Invalid page numbers")
    doc.select(sorted(set(pages)))
    return _out(doc)


def merge(files: list[bytes]) -> bytes:
    out = fitz.open()
    for f in files:
        d = _open(f)
        out.insert_pdf(d)
        d.close()
    return _out(out)


def parse_ranges(spec: str, n: int) -> list[list[int]]:
    """'1-3,5,7-9' -> [[0,1,2],[4],[6,7,8]] (zero-based)."""
    groups = []
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        try:
            if "-" in part:
                a, b = part.split("-", 1)
                a, b = int(a), int(b) if b else n
            else:
                a = b = int(part)
        except ValueError:
            raise HTTPException(400, f"Bad range '{part}'")
        if not (1 <= a <= b <= n):
            raise HTTPException(400, f"Range '{part}' is outside 1-{n}")
        groups.append(list(range(a - 1, b)))
    if not groups:
        raise HTTPException(400, "No ranges given")
    return groups


def split(data: bytes, groups: list[list[int]], stem: str) -> bytes:
    """Returns a ZIP with one PDF per group."""
    src = _open(data)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as z:
        for i, g in enumerate(groups, 1):
            d = fitz.open()
            for p in g:
                d.insert_pdf(src, from_page=p, to_page=p)
            label = f"{g[0] + 1}" if len(g) == 1 else f"{g[0] + 1}-{g[-1] + 1}"
            z.writestr(f"{stem}_part{i}_p{label}.pdf", _out(d))
    return buf.getvalue()


# ---------------------------------------------------------------- compress / OCR
_LEVELS = {
    "low": (200, 150, 85),
    "medium": (150, 110, 70),
    "high": (120, 80, 50),
    # A lower threshold deliberately catches images already reduced by High (80 dpi).
    "maximum": (70, 48, 30),
}


def check_level(level: str):
    if level not in _LEVELS:
        raise HTTPException(400, "level must be low, medium, high or maximum")


def compress_file(src, dst, level: str) -> None:
    """Long-running: called from a background job. Reads src from disk, writes the smaller copy to dst."""
    check_level(level)
    doc = _open(src)
    thr, tgt, q = _LEVELS[level]
    try:
        doc.rewrite_images(dpi_threshold=thr, dpi_target=tgt, quality=q)
    except Exception:
        pass  # older PyMuPDF / odd images: structural compression below still applies
    doc.save(str(dst), garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True)
    doc.close()


def ocr_available() -> bool:
    return shutil.which("ocrmypdf") is not None and shutil.which("tesseract") is not None


class Cancelled(Exception):
    pass


def _kill_tree(p: subprocess.Popen):
    try:
        if os.name == "nt":
            subprocess.run(["taskkill", "/PID", str(p.pid), "/T", "/F"], capture_output=True)
        else:
            import signal
            os.killpg(os.getpgid(p.pid), signal.SIGTERM)
            try:
                p.wait(10)
            except subprocess.TimeoutExpired:
                os.killpg(os.getpgid(p.pid), signal.SIGKILL)
    except (ProcessLookupError, OSError):
        pass
    try:
        p.wait(15)
    except subprocess.TimeoutExpired:
        p.kill()


def check_ocr_args(langs: str | None, pages: str | None) -> tuple[str, str | None]:
    if not ocr_available():
        raise HTTPException(501, "OCR engine (ocrmypdf + tesseract) is not installed on the server")
    langs = (langs or config.OCR_LANGS).strip()
    if not langs or not all(c.isalnum() or c in "+_" for c in langs):
        raise HTTPException(400, "Bad language code")
    pages = (pages or "").strip().replace(" ", "") or None
    if pages and not re.fullmatch(r"\d+(-\d+)?(,\d+(-\d+)?)*", pages):
        raise HTTPException(400, "Pages must look like 1-50 or 1,3,7-9")
    return langs, pages


def ocr_file(src, dst, langs: str, redo: bool, deskew: bool, pages: str | None, cancelled, log_path) -> None:
    """Scanned PDF -> searchable PDF using OCRmyPDF (Tesseract). Runs as a job; polls `cancelled()` so it can be stopped.
    Existing OCR text is replaced only when requested; visible page artwork is preserved."""
    cmd = ["ocrmypdf", "-l", langs, "--output-type", "pdf", "--optimize", "0", "--jobs", str(config.OCR_JOBS)]
    progress_plugin = Path(os.getenv("PDFDESK_OCR_PROGRESS_PLUGIN", Path(__file__).with_name("ocr_progress_plugin.py")))
    if progress_plugin.is_file():
        cmd += ["--plugin", str(progress_plugin)]
    cmd += ["--redo-ocr"] if redo else ["--skip-text"]
    if deskew:
        cmd += ["--deskew"]
    if pages:
        cmd += ["--pages", pages]
    cmd += [str(src), str(dst)]
    progress_file = str(Path(log_path).with_suffix(".progress.json"))
    try:
        progress_doc = fitz.open(str(src), filetype="pdf")
        total_pages = progress_doc.page_count
        progress_doc.close()
    except Exception:
        total_pages = 0
    env = {**os.environ, "TMPDIR": str(config.TMP_DIR), "TEMP": str(config.TMP_DIR), "TMP": str(config.TMP_DIR),
           "PDFDESK_OCR_PROGRESS_FILE": progress_file, "PDFDESK_OCR_TOTAL_PAGES": str(total_pages)}
    with open(log_path, "wb") as lf:
        # own process group, so cancelling stops OCRmyPDF *and* the Tesseract/Ghostscript workers it spawned
        popen_options = {"creationflags": subprocess.CREATE_NO_WINDOW} if os.name == "nt" else {"start_new_session": True}
        p = subprocess.Popen(cmd, stdout=lf, stderr=lf, env=env, **popen_options)
        while p.poll() is None:
            if cancelled():
                _kill_tree(p)
                raise Cancelled()
            time.sleep(1)
    if p.returncode != 0 or not Path(dst).exists():
        tail = Path(log_path).read_text(errors="replace")[-400:]
        raise HTTPException(500, "OCR failed: " + tail)


def ocr_progress(log_path) -> int | None:
    """Best effort: the last 'NN%' OCRmyPDF printed to its log."""
    try:
        m = re.findall(r"(\d{1,3})%", Path(log_path).read_text(errors="replace")[-4000:])
        return min(100, int(m[-1])) if m else None
    except Exception:
        return None


def ocr_progress_detail(log_path) -> dict:
    """Read machine-readable phase and page completion emitted by the OCRmyPDF plugin."""
    progress_path = Path(log_path).with_suffix(".progress.json")
    try:
        state = json.loads(progress_path.read_text(encoding="utf-8"))
        total = state.get("total")
        completed = state.get("completed", 0)
        if "completed" in state and isinstance(total, (int, float)) and total > 0:
            state["progress"] = min(100, max(0, round(completed / total * 100)))
        else:
            state["progress"] = None
        if state.get("stage", "").strip().lower() == "ocr" and state.get("active_page"):
            state["progress_label"] = f"OCR scanning page {int(state['active_page'])} of {int(total or 0)}"
        elif state.get("stage", "").strip().lower() == "ocr" and state.get("last_page"):
            state["progress_label"] = f"OCR finished page {int(state['last_page'])}; preparing the next page…"
        else:
            state["progress_label"] = state.get("stage") or "Preparing OCR…"
        return state
    except (OSError, ValueError, TypeError):
        return {}


def outline(data: bytes) -> list[dict]:
    doc = _open(data)
    out = [{"level": lvl, "title": t, "page": max(0, p - 1)} for lvl, t, p in doc.get_toc(simple=True) if p > 0]
    _release(doc)
    return out


# ---------------------------------------------------------------- search / export
SEARCH_LIMIT = 2000


def _numeric_ocr_char_hits(rawdict: dict, query: str) -> list[tuple[fitz.Rect, str]]:
    """Find numeric matches from one raw text extraction, including split OCR glyphs."""
    if not re.fullmatch(r"[\d\s\-()./]+", query.strip()):
        return []
    target = re.sub(r"\D", "", query)
    if len(target) < 3:
        return []

    confusables = str.maketrans({"O": "0", "o": "0", "I": "1", "l": "1", "|": "1",
                                "S": "5", "s": "5", "B": "8", "Z": "2", "z": "2", "G": "6"})
    rows = []
    for block in rawdict.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            chars = [ch for span in line.get("spans", []) for ch in span.get("chars", [])]
            if not chars:
                continue
            y0, y1 = line["bbox"][1], line["bbox"][3]
            center = (y0 + y1) / 2
            row = next((r for r in rows if abs(r[0] - center) <= max(2.0, min(r[1] - r[2], y1 - y0) * 0.4)), None)
            if row is None:
                row = [center, y1, y0, []]
                rows.append(row)
            else:
                row[1] = max(row[1], y1)
                row[2] = min(row[2], y0)
            row[3].extend(chars)

    hits = []
    for row in rows:
        chars = sorted(row[3], key=lambda ch: (ch["bbox"][0], ch["bbox"][1]))
        cluster = []
        previous = None

        def flush():
            if not cluster:
                return
            joined = "".join(ch[0] for ch in cluster)
            offset = joined.find(target)
            while offset >= 0:
                matched = [ch[1] for ch in cluster[offset:offset + len(target)]]
                rect = fitz.Rect(min(b[0] for b in matched), min(b[1] for b in matched),
                                 max(b[2] for b in matched), max(b[3] for b in matched))
                snippet = "".join(ch.get("c", "") for ch in chars)[:200]
                hits.append((rect, snippet))
                offset = joined.find(target, offset + 1)

        for ch in chars:
            value = ch.get("c", "")
            x0, y0, x1, y1 = ch["bbox"]
            if value.isspace():
                if previous is None or x0 - previous[2] > max(2.0, min(y1 - y0, previous[3] - previous[1]) * 2.0):
                    flush()
                    cluster.clear()
                    previous = None
                continue
            normalized = value.translate(confusables)
            if normalized not in "0123456789":
                flush()
                cluster.clear()
                previous = None
                continue
            if previous is not None:
                gap = x0 - previous[2]
                max_gap = max(2.0, min(y1 - y0, previous[3] - previous[1]) * 2.0)
                if gap > max_gap:
                    flush()
                    cluster.clear()
            cluster.append((normalized, (x0, y0, x1, y1)))
            previous = (x0, y0, x1, y1)
        flush()
    return hits


def search_text(data, query: str, start: int = 0, stop: int | None = None) -> list[dict]:
    """Stops at SEARCH_LIMIT matches so a common word in a 900-page file doesn't flood the browser.
    start/stop let the caller search a huge file in slices, so other requests aren't blocked meanwhile."""
    doc = _open(data)
    hits = []
    numeric_query = bool(re.fullmatch(r"[\d\s\-()./]+", query.strip()) and len(re.sub(r"\D", "", query)) >= 3)
    for i in range(start, doc.page_count if stop is None else min(stop, doc.page_count)):
        pg = doc[i]
        if numeric_query:
            # Numeric OCR searches need glyph-level recovery when scanners split
            # digits into spans or confuse similar characters. Do one extraction
            # per page; search_for + words + rawdict was doing up to three.
            for rect, snippet in _numeric_ocr_char_hits(pg.get_text("rawdict"), query):
                hits.append({"page": i, "rect": _to_view(pg, rect), "snippet": snippet})
            if len(hits) >= SEARCH_LIMIT:
                hits = hits[:SEARCH_LIMIT]
                break
            continue
        found = pg.search_for(query)
        if not found:
            continue
        words, lines = [], {}
        if found:
            words = pg.get_text("words", textpage=pg.get_textpage())  # (x0,y0,x1,y1,text,block,line,word_no)
            for w in words:
                lines.setdefault((w[5], w[6]), []).append(w)
        for r in found:
            # a few words either side of the match on its own line, so the list shows *where* it is, not just the page
            cx, cy = (r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2
            hw = next((w for w in words if w[0] - 1 <= cx <= w[2] + 1 and w[1] - 1 <= cy <= w[3] + 1), None)
            snippet = ""
            if hw:
                lw = sorted(lines[(hw[5], hw[6])], key=lambda w: w[7])
                k = lw.index(hw)
                snippet = " ".join(w[4] for w in lw[max(0, k - 7): k + 10])[:200]
            hits.append({"page": i, "rect": _to_view(pg, r), "snippet": snippet})
        if len(hits) >= SEARCH_LIMIT:
            hits = hits[:SEARCH_LIMIT]
            break
    _release(doc)
    return hits


def to_text(data: bytes) -> str:
    doc = _open(data)
    t = "\n\f".join(p.get_text() for p in doc)
    doc.close()
    return t


def export_images(data: bytes, fmt: str, dpi: int, stem: str) -> bytes:
    if fmt not in ("png", "jpg"):
        raise HTTPException(400, "fmt must be png or jpg")
    doc = _open(data)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as z:
        for i, pg in enumerate(doc, 1):
            pix = pg.get_pixmap(dpi=max(50, min(dpi, 300)), alpha=False)
            z.writestr(f"{stem}_page{i:03d}.{fmt}", pix.tobytes("jpeg" if fmt == "jpg" else "png"))
    return buf.getvalue()


def export_html(data: bytes) -> str:
    doc = _open(data)
    body = "\n".join(f'<div class="page" id="p{i + 1}">{p.get_text("html")}</div>' for i, p in enumerate(doc))
    doc.close()
    return f'<!doctype html><meta charset="utf-8"><title>Exported PDF</title>{body}'


def export_docx(data: bytes) -> bytes:
    """Text and images per page into a Word file (layout is simplified, not pixel-exact)."""
    from docx import Document
    from docx.shared import Inches

    doc = _open(data)
    out = Document()
    for i, pg in enumerate(doc):
        if i:
            out.add_page_break()
        for b in pg.get_text("dict", sort=True)["blocks"]:
            if b.get("type") == 0:
                text = "\n".join("".join(s["text"] for s in ln["spans"]) for ln in b["lines"]).strip()
                if text:
                    out.add_paragraph(text)
            elif b.get("type") == 1 and b.get("image"):
                try:
                    out.add_picture(io.BytesIO(b["image"]), width=Inches(min(6, (b["bbox"][2] - b["bbox"][0]) / 72)))
                except Exception:
                    pass
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def export_xlsx(data: bytes) -> bytes:
    """Export extracted PDF text as Excel rows without requiring an external spreadsheet package."""
    doc = _open(data)
    rows = [["Page", "Block", "Y position (pt)", "Text"]]
    for page_no, page in enumerate(doc, 1):
        for block_no, block in enumerate(page.get_text("blocks", sort=True), 1):
            text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(block[4])).strip()
            if text:
                rows.append([page_no, block_no, round(float(block[1]), 1), text])
    doc.close()

    # Excel allows 1,048,576 rows per sheet. Repeat the header if a document needs more than one sheet.
    max_data_rows = 1_000_000
    chunks = [rows[1 + i:1 + i + max_data_rows] for i in range(0, max(1, len(rows) - 1), max_data_rows)]
    sheet_rows = [[rows[0], *chunk] for chunk in chunks]

    def col_name(index: int) -> str:
        name = ""
        while index:
            index, rem = divmod(index - 1, 26)
            name = chr(65 + rem) + name
        return name

    def cell(ref: str, value: Any) -> str:
        if isinstance(value, (int, float)):
            return f'<c r="{ref}"><v>{value}</v></c>'
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))
        preserve = ' xml:space="preserve"' if value[:1].isspace() or value[-1:].isspace() else ""
        # Inline strings keep PDF content as text, including values beginning with '=' (never formulas).
        return f'<c r="{ref}" t="inlineStr"><is><t{preserve}>{escape(value)}</t></is></c>'

    def worksheet(sheet: list[list[Any]]) -> str:
        body = []
        for r, values in enumerate(sheet, 1):
            body.append(f'<row r="{r}">' + "".join(cell(f"{col_name(c)}{r}", v) for c, v in enumerate(values, 1)) + "</row>")
        return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                '<sheetData>' + "".join(body) + '</sheetData></worksheet>')

    sheets = "".join(f'<sheet name="PDF Text {i}" sheetId="{i}" r:id="rId{i}"/>'
                     for i in range(1, len(sheet_rows) + 1))
    relationships = "".join(
        f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>'
        for i in range(1, len(sheet_rows) + 1)
    )
    content_types = ''.join(
        f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        for i in range(1, len(sheet_rows) + 1)
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                   + content_types + '</Types>')
        z.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
                   '</Relationships>')
        z.writestr("xl/workbook.xml", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                   'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'
                   + sheets + '</sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   + relationships + '</Relationships>')
        for i, sheet in enumerate(sheet_rows, 1):
            z.writestr(f"xl/worksheets/sheet{i}.xml", worksheet(sheet))
    return out.getvalue()


# ---------------------------------------------------------------- create
_SIZES = {"a4": (595, 842), "letter": (612, 792), "legal": (612, 1008), "a3": (842, 1191)}


def create_blank(count: int, size: str, landscape: bool = False) -> bytes:
    w, h = _SIZES.get(size.lower(), _SIZES["a4"])
    if landscape:
        w, h = h, w
    d = fitz.open()
    for _ in range(max(1, min(count, 500))):
        d.new_page(width=w, height=h)
    return _out(d)


def create_from_text(text: str, size: str, fontsize: float) -> bytes:
    w, h = _SIZES.get(size.lower(), _SIZES["a4"])
    fontsize = max(6, min(fontsize, 48))
    margin, lead = 54, fontsize * 1.35
    maxw = w - 2 * margin
    lines = []
    for para in text.replace("\r", "").split("\n"):
        words, cur = para.split(" "), ""
        for wd in words:
            t = (cur + " " + wd).strip() if cur else wd
            if fitz.get_text_length(t, "helv", fontsize) <= maxw:
                cur = t
            else:
                lines.append(cur)
                cur = wd
        lines.append(cur)
    per = max(1, int((h - 2 * margin) // lead))
    d = fitz.open()
    for s in range(0, max(len(lines), 1), per):
        pg = d.new_page(width=w, height=h)
        y = margin + fontsize
        for ln in lines[s:s + per]:
            pg.insert_text((margin, y), ln, fontsize=fontsize, fontname="helv")
            y += lead
    return _out(d)


def images_to_pdf(images: list[bytes]) -> bytes:
    out = fitz.open()
    for im in images:
        try:
            pix = fitz.open(stream=im)
            pdf = fitz.open("pdf", pix.convert_to_pdf())
        except Exception:
            raise HTTPException(400, "Unsupported image file")
        out.insert_pdf(pdf)
    return _out(out)


# ---------------------------------------------------------------- forms
_WTYPES = {
    "text": fitz.PDF_WIDGET_TYPE_TEXT,
    "checkbox": fitz.PDF_WIDGET_TYPE_CHECKBOX,
    "dropdown": fitz.PDF_WIDGET_TYPE_COMBOBOX,
    "signature": fitz.PDF_WIDGET_TYPE_SIGNATURE,
}


def list_fields(data) -> list[dict]:
    doc = _open(data)
    out = []
    if not doc.is_form_pdf:  # no form at all: skip scanning every page
        _release(doc)
        return out
    for i in range(doc.page_count):
        pg = doc[i]
        for w in pg.widgets() or []:
            kind = w.field_type_string
            val = w.field_value
            if kind == "CheckBox":
                val = val not in (False, "Off", "", None, 0)
            out.append({"page": i, "name": w.field_name, "type": kind, "rect": _to_view(pg, w.rect),
                        "value": val, "choices": list(w.choice_values or []), "readonly": bool(w.field_flags & 1)})
    _release(doc)
    return out


def fill_fields(data: bytes, values: dict) -> bytes:
    doc = _open(data)
    done = 0
    for pg in doc:
        for w in pg.widgets() or []:
            if w.field_name in values and w.field_type_string != "Signature":
                v = values[w.field_name]
                w.field_value = bool(v) if w.field_type_string == "CheckBox" else str(v)
                w.update()
                done += 1
    if not done:
        raise HTTPException(400, "No matching form fields")
    return _out(doc)


def add_field(data: bytes, page: int, rect, kind: str, name: str, options: list[str] | None = None) -> bytes:
    if kind not in _WTYPES:
        raise HTTPException(400, "Unknown field type")
    doc = _open(data)
    pg = _page(doc, page)
    existing = {w.field_name for p in doc for w in (p.widgets() or [])}
    name = (name or kind).strip()[:60] or kind
    base, k = name, 2
    while name in existing:
        name, k = f"{base}_{k}", k + 1
    w = fitz.Widget()
    w.field_type = _WTYPES[kind]
    w.field_name = name
    w.rect = _to_page(pg, rect)
    w.border_color = (0.2, 0.3, 0.9)
    w.border_width = 1
    w.fill_color = (0.93, 0.95, 1)
    if kind == "text":
        w.text_fontsize = 0  # auto
        w.field_value = ""
    elif kind == "dropdown":
        w.choice_values = [o for o in (options or []) if o.strip()] or ["Option 1", "Option 2"]
        w.field_value = w.choice_values[0]
    pg.add_widget(w)
    return _out(doc)


# ---------------------------------------------------------------- e-signature / stamps
def _data_url(url: str) -> bytes:
    try:
        return base64.b64decode(url.split(",", 1)[1])
    except Exception:
        raise HTTPException(400, "Bad image data")


def place_image(data: bytes, page: int, rect, image_data_url: str, label: str = "Image", author: str = "") -> bytes:
    """Signatures / pasted pictures become Stamp annotations, so they can be selected, moved, resized and deleted."""
    doc = _open(data)
    pg = _page(doc, page)
    ann = pg.add_stamp_annot(_to_page(pg, rect), stamp=_data_url(image_data_url))
    ann.set_info(title=author[:60], content=label[:60])
    ann.update()
    return _out(doc)


STAMPS = {
    "APPROVED": (0, 0.5, 0), "REJECTED": (0.75, 0, 0), "DRAFT": (0.4, 0.4, 0.4), "CONFIDENTIAL": (0.75, 0, 0),
    "FOR REVIEW": (0, 0.3, 0.8), "PAID": (0, 0.5, 0), "RECEIVED": (0, 0.3, 0.8), "VOID": (0.75, 0, 0),
    "COPY": (0.4, 0.4, 0.4), "COMPLETED": (0, 0.5, 0),
}


def stamp(data: bytes, pages: list[int], center, text: str, color, size: float, with_date: bool, by: str) -> bytes:
    """Stamps are real Stamp annotations (an image of the stamp), so they can be moved, resized and deleted."""
    text = text.strip()[:40].upper() or "APPROVED"
    size = max(10, min(size, 60))
    col = _rgb(color)
    sub = f"{date.today().isoformat()}" + (f"  {by}" if by else "") if with_date else ""
    w = max(fitz.get_text_length(text, "hebo", size), fitz.get_text_length(sub, "helv", size * 0.45)) + size
    h = size * 1.5 + (size * 0.7 if sub else 0)
    tmp = fitz.open()
    tp = tmp.new_page(width=w, height=h)
    r = fitz.Rect(2, 2, w - 2, h - 2)
    tp.draw_rect(r, color=col, width=2.2, radius=0.12)
    tp.insert_textbox(fitz.Rect(r.x0, r.y0 + size * 0.15, r.x1, r.y0 + size * 1.5), text, fontsize=size,
                      fontname="hebo", color=col, align=1)
    if sub:
        tp.insert_textbox(fitz.Rect(r.x0, r.y0 + size * 1.45, r.x1, r.y1), sub, fontsize=size * 0.45,
                          fontname="helv", color=col, align=1)
    png = tp.get_pixmap(matrix=fitz.Matrix(4, 4), alpha=True).tobytes("png")
    doc = _open(data)
    for p in pages:
        pg = _page(doc, p)
        c = _pt_to_page(pg, center)
        ann = pg.add_stamp_annot(fitz.Rect(c.x - w / 2, c.y - h / 2, c.x + w / 2, c.y + h / 2), stamp=png)
        ann.set_info(title=by[:60], content=text)
        ann.update()
    return _out(doc)


# ---------------------------------------------------------------- edit existing objects
def _find_annot(pg, xref: int):
    for an in pg.annots() or []:
        if an.xref == xref:
            return an
    raise HTTPException(404, "Object not found")


def update_annot(data: bytes, page: int, xref: int, rect=None, text=None, color=None, width=None, size=None) -> bytes:
    doc = _open(data)
    pg = _page(doc, page)
    an = _find_annot(pg, xref)
    kind = an.type[1]
    if rect is not None:
        if kind not in ("Stamp", "FreeText", "Square", "Circle", "Text"):
            raise HTTPException(400, f"A {kind} object can't be moved or resized — delete it and add it again")
        an.set_rect(_to_page(pg, rect))
    if text is not None:
        an.set_info(content=text)
    if color is not None:
        c = _rgb(color)
        if kind == "FreeText":
            an.update(text_color=c)
        elif kind in ("Square", "Circle", "Ink", "Highlight", "Underline", "StrikeOut"):
            an.set_colors(stroke=c)
    if width is not None and kind in ("Square", "Circle", "Ink"):
        an.set_border(width=max(0.5, min(float(width), 20)))
    if kind == "FreeText" and size is not None:
        an.update(fontsize=max(5, min(float(size), 96)))
    else:
        an.update()
    return _out(doc)


def delete_field(data: bytes, name: str) -> bytes:
    doc = _open(data)
    for pg in doc:
        for w in pg.widgets() or []:
            if w.field_name == name:
                pg.delete_widget(w)
                return _out(doc)
    raise HTTPException(404, "Field not found")


def update_field(data: bytes, name: str, rect=None, new_name: str | None = None) -> bytes:
    doc = _open(data)
    for pg in doc:
        for w in pg.widgets() or []:
            if w.field_name == name:
                if rect is not None:
                    w.rect = _to_page(pg, rect)
                if new_name and new_name != name:
                    if any(x.field_name == new_name for p in doc for x in (p.widgets() or [])):
                        raise HTTPException(400, "A field with that name already exists")
                    w.update()
                    safe = new_name[:60].replace("\\", "").replace("(", "").replace(")", "")
                    parent = doc.xref_get_key(w.xref, "Parent")
                    target = int(parent[1].split()[0]) if parent[0] == "xref" else w.xref
                    doc.xref_set_key(target, "T", f"({safe})")
                    return _out(doc)
                w.update()
                return _out(doc)
    raise HTTPException(404, "Field not found")


def watermark(data: bytes, text: str) -> bytes:
    doc = _open(data)
    for pg in doc:
        r = pg.rect
        pg.insert_text(
            fitz.Point(r.width * 0.15, r.height * 0.55), text, fontsize=r.width / max(len(text), 6) * 1.6,
            fontname="helv", color=(0.8, 0.8, 0.8), overlay=True, fill_opacity=0.5,
        )
    return _out(doc)


# ---------------------------------------------------------------- protect
def protect(data: bytes, user_pw: str, owner_pw: str, allow_print: bool, allow_copy: bool, allow_edit: bool) -> bytes:
    if not user_pw and not owner_pw:
        raise HTTPException(400, "Enter an open password or an owner (permissions) password")
    perm = fitz.PDF_PERM_ACCESSIBILITY
    if allow_print:
        perm |= fitz.PDF_PERM_PRINT | fitz.PDF_PERM_PRINT_HQ
    if allow_copy:
        perm |= fitz.PDF_PERM_COPY
    if allow_edit:
        perm |= fitz.PDF_PERM_MODIFY | fitz.PDF_PERM_ANNOTATE | fitz.PDF_PERM_FORM | fitz.PDF_PERM_ASSEMBLE
    doc = _open(data)
    out = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=user_pw or "",
                      owner_pw=owner_pw or secrets.token_hex(16), permissions=perm, garbage=3, deflate=True)
    doc.close()
    return out


# ---------------------------------------------------------------- compare
def _page_lines(doc: fitz.Document, i: int) -> list[str]:
    if i >= doc.page_count:
        return []
    return [re.sub(r"\s+", " ", ln).strip() for ln in doc[i].get_text().splitlines() if ln.strip()]


def compare(a: bytes, b: bytes) -> dict:
    da, db = _open(a), _open(b)
    changes, added, removed = [], 0, 0
    for i in range(max(da.page_count, db.page_count)):
        la, lb = _page_lines(da, i), _page_lines(db, i)
        for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, la, lb, autojunk=False).get_opcodes():
            if tag == "equal":
                continue
            changes.append({"page": i + 1, "removed": la[i1:i2], "added": lb[j1:j2]})
            removed += i2 - i1
            added += j2 - j1
    res = {"pages_a": da.page_count, "pages_b": db.page_count, "changes": changes[:2000],
           "added": added, "removed": removed, "identical": not changes,
           "no_text": not any(_page_lines(da, i) for i in range(da.page_count)) or
                      not any(_page_lines(db, i) for i in range(db.page_count))}
    da.close()
    db.close()
    return res


# ---------------------------------------------------------------- summary
_STOP = set("""a an the and or but if then else of to in on at by for with from as is are was were be been being it its this
that these those i you he she we they them his her our your their not no yes do does did have has had will would can could
should may might must also than so such into over under about after before between which who whom whose what when where why how
all any each more most other some very just only""".split())


def extractive_summary(text: str, sentences: int) -> list[str]:
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)  # PDF hard line breaks are not sentence breaks
    seen, sents = set(), []
    for s in re.split(r"(?<=[.!?])\s+|\n{2,}", re.sub(r"[ \t]+", " ", text)):
        s = s.strip()
        if len(s.split()) >= 5 and s.lower() not in seen:
            seen.add(s.lower())
            sents.append(s)
    if len(sents) <= sentences:
        return sents
    freq = Counter(w for w in re.findall(r"[A-Za-z][A-Za-z'-]{2,}", text.lower()) if w not in _STOP)
    top = max(freq.values(), default=1)
    scored = []
    for i, s in enumerate(sents):
        ws = [w for w in re.findall(r"[A-Za-z][A-Za-z'-]{2,}", s.lower()) if w not in _STOP]
        score = sum(freq[w] / top for w in ws) / (len(ws) ** 0.6 or 1)
        scored.append((score + (0.15 if i < 3 else 0), i))
    keep = sorted(i for _, i in sorted(scored, reverse=True)[:sentences])
    return [sents[i] for i in keep]


def llm_summary(text: str, length: str) -> str:
    import httpx

    words = {"short": "about 80 words", "medium": "about 200 words", "long": "about 450 words"}[length]
    r = httpx.post(
        "https://api.anthropic.com/v1/messages",
        headers={"x-api-key": config.ANTHROPIC_API_KEY, "anthropic-version": "2023-06-01"},
        json={"model": config.LLM_MODEL, "max_tokens": 1500,
              "messages": [{"role": "user", "content":
                            f"Summarise the following document in {words}. Start with one sentence on what the document is, "
                            f"then give the key points as short bullets (facts, dates, amounts, decisions, action items). "
                            f"Do not invent anything that is not in the text.\n\n<document>\n{text[:150000]}\n</document>"}]},
        timeout=120,
    )
    if r.status_code != 200:
        raise HTTPException(502, f"Summary service error ({r.status_code})")
    return "".join(b.get("text", "") for b in r.json().get("content", [])).strip()


def summarize(data: bytes, length: str) -> dict:
    if length not in ("short", "medium", "long"):
        raise HTTPException(400, "length must be short, medium or long")
    text = to_text(data)
    if len(text.split()) < 5:
        raise HTTPException(400, "This PDF has no readable text. Run Scan & OCR first.")
    base = {"words": len(text.split())}
    if config.ANTHROPIC_API_KEY:
        try:
            return {**base, "engine": "claude", "summary": llm_summary(text, length)}
        except Exception as exc:
            base["notice"] = f"AI summary unavailable; showing an offline summary instead ({exc.detail if isinstance(exc, HTTPException) else 'service connection failed'})."
    n = {"short": 3, "medium": 6, "long": 12}[length]
    summary = "\n".join("• " + s for s in extractive_summary(text, n))
    if not summary:
        raise HTTPException(400, "Could not find readable sentences. Run Scan & OCR first.")
    return {**base, "engine": "extractive", "summary": summary}
